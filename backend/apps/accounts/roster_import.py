import csv
from pathlib import Path

from django.db import transaction

from apps.accounts.models import User
from apps.accounts.student_email import student_email_from_mc
from apps.common.choices import AccountStatus, UserRole
from apps.common.validators import validate_mc_number


def _normalize_header(name: str) -> str:
    return name.strip().lower().replace(" ", "_")


def iter_roster_rows(path: Path):
    suffix = path.suffix.lower()
    if suffix == ".csv":
        with path.open(newline="", encoding="utf-8-sig") as handle:
            reader = csv.DictReader(handle)
            if not reader.fieldnames:
                raise ValueError("The roster file has no header row.")
            field_map = {_normalize_header(name): name for name in reader.fieldnames}
            mc_key = field_map.get("mc_number")
            cpm_key = field_map.get("cpm_number")
            if not mc_key or not cpm_key:
                raise ValueError("CSV must include mc_number and cpm_number columns.")
            for row in reader:
                yield str(row.get(mc_key) or "").strip(), str(row.get(cpm_key) or "").strip()
        return

    if suffix in {".xlsx", ".xlsm"}:
        try:
            import openpyxl
        except ImportError as exc:
            raise ValueError("Install openpyxl to import Excel roster files.") from exc

        workbook = openpyxl.load_workbook(path, data_only=True, read_only=True)
        sheet_name = "IT" if "IT" in workbook.sheetnames else workbook.sheetnames[0]
        worksheet = workbook[sheet_name]
        rows = worksheet.iter_rows(values_only=True)
        header = next(rows, None)
        if not header:
            raise ValueError("The roster worksheet has no header row.")
        field_map = {_normalize_header(str(cell or "")): index for index, cell in enumerate(header)}
        cpm_index = field_map.get("cpm_number")
        mc_index = field_map.get("mc_number")
        if cpm_index is None or mc_index is None:
            raise ValueError("Worksheet must include Cpm Number and Mc Number columns.")
        for row in rows:
            if not row:
                continue
            cpm = str(row[cpm_index] or "").strip()
            mc = str(row[mc_index] or "").strip()
            if mc and cpm:
                yield mc, cpm
        workbook.close()
        return

    raise ValueError(f"Unsupported roster file type: {path.suffix}")


def import_student_roster_file(path: Path) -> tuple[int, int]:
    created = 0
    skipped = 0
    with transaction.atomic():
        for raw_mc, raw_cpm in iter_roster_rows(path):
            if not raw_cpm.strip():
                raise ValueError("Every student row must include a non-empty CPM number.")
            try:
                mc_number = validate_mc_number(raw_mc)
            except ValueError as exc:
                raise ValueError(f"Invalid MC number '{raw_mc}': {exc}") from exc

            email = student_email_from_mc(mc_number)
            existing = User.objects.filter(mc_number=mc_number, role=UserRole.STUDENT).first()
            if existing:
                if existing.email.lower() != email.lower():
                    if User.objects.filter(email__iexact=email).exclude(pk=existing.pk).exists():
                        raise ValueError(
                            f"Cannot update student {mc_number}: email {email} already exists."
                        )
                    existing.email = email
                    existing.save(update_fields=["email", "updated_at"])
                skipped += 1
                continue
            if User.objects.filter(email__iexact=email).exists():
                raise ValueError(f"Cannot create student {mc_number}: email {email} already exists.")

            User.objects.create_user(
                email=email,
                password=raw_cpm,
                name=f"Student {mc_number}",
                mc_number=mc_number,
                role=UserRole.STUDENT,
                status=AccountStatus.ACTIVE,
            )
            created += 1

    return created, skipped
